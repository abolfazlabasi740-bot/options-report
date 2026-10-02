#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Audit OptionSchool24 as a supplemental-only evidence source.

Never edits TSETMC snapshots, identities, scores, ranking, or production gates.
OptionSchool rows remain source-local until an independently verified mapping exists.
"""
from __future__ import annotations
import hashlib, json, math
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook

SOURCE = "OPTIONSCHOOL24_SUPPLEMENTAL_ONLY"
FILE = Path("data/optionschool24_live.xlsx")
OUT = Path("data/optionschool24_supplemental_audit.json")

def clean(v):
    if v is None: return None
    if isinstance(v, str):
        v = v.strip().replace(",", "")
        if not v: return None
    try:
        x = float(v)
        return x if math.isfinite(x) else None
    except (TypeError, ValueError):
        return str(v).strip() or None

def main():
    raw = FILE.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    wb = load_workbook(FILE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    headers = [str(x).strip() if x is not None else "" for x in next(it)]
    rows = [dict(zip(headers, r)) for r in it if any(x is not None for x in r)]
    required_relative = ["نماد", "قیمت اعمال", "قیمت سهم پایه", "تاریخ سررسید",
                         "آخرین قیمت", "اختلاف تا بلک شولز", "بلک شولز", "نوسان ضمنی"]
    relative_missing = {k: sum(1 for r in rows if clean(r.get(k)) is None) for k in required_relative}
    relative_complete = sum(1 for r in rows if all(clean(r.get(k)) is not None for k in required_relative))
    chain_required = ["نماد", "قیمت اعمال", "تاریخ سررسید"]
    chain_missing = {k: sum(1 for r in rows if clean(r.get(k)) is None) for k in chain_required}
    # The workbook has no explicit underlying identifier/symbol and no explicit CALL/PUT field.
    # Do not infer either from option symbol or underlying price.
    payload = {
      "status": "AUDIT_COMPLETE",
      "source": SOURCE,
      "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
      "file": str(FILE), "file_bytes": len(raw), "sha256": sha,
      "sheet": ws.title, "data_rows": len(rows), "column_count": len(headers),
      "columns": headers,
      "relative_value_anomaly": {
        "status": "RAW_FIELDS_AVAILABLE_REQUIRES_MAPPING_AND_VALIDATION",
        "required_fields": required_relative,
        "missing_rows_by_field": relative_missing,
        "rows_with_all_raw_fields": relative_complete,
        "available_raw_fields": ["بلک شولز", "اختلاف تا بلک شولز", "نوسان ضمنی", "نوسان تاریخی",
          "ارزش ذاتی", "ارزش زمانی", "سر به سر", "اختلاف تا سر به سر",
          "آخرین قیمت", "قیمت سهم پایه", "قیمت اعمال", "حجم معاملات", "ارزش معاملات",
          "قیمت بهترین تقاضا", "قیمت بهترین عرضه", "دلتا", "تتا", "گاما", "وگا", "رو"],
        "constraints": [
          "OptionSchool rows are source-local; no TSETMC instrument_id is present in this workbook.",
          "Do not promote symbol-only matches to exact instrument identity.",
          "Do not change TSETMC snapshots, scores, ranking, or source-of-truth.",
          "Raw Black-Scholes/IV availability is not itself independent economic validation."
        ]
      },
      "chain_structure_anomaly": {
        "status": "IDENTITY_FIELDS_MISSING",
        "missing_explicit_fields": ["underlying_id_or_symbol", "contract_type_CALL_PUT"],
        "available_fields": ["نماد", "قیمت اعمال", "تاریخ سررسید", "قیمت سهم پایه"],
        "missing_rows_by_available_field": chain_missing,
        "reason": "Underlying price is not an underlying identity; option symbol is not parsed to infer underlying or CALL/PUT.",
        "constraints": ["No guessed chain grouping.", "No member-score dispersion computed from OptionSchool raw fields."]
      },
      "project_boundary": {
        "tsetmc_modified": False,
        "tsetmc_snapshot_files_modified": False,
        "supplemental_only": True,
        "production_signal_authorization": "FORBIDDEN",
        "g7_5_closure": "NOT_AUTHORIZED_BY_THIS_AUDIT_ALONE"
      }
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: payload[k] for k in ("status","file_bytes","sha256","data_rows","column_count","relative_value_anomaly","chain_structure_anomaly","project_boundary")}, ensure_ascii=False, indent=2))

if __name__ == "__main__": main()
